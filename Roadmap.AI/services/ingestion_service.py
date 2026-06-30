from io import BytesIO
from typing import cast

import fitz
from azure.storage.blob import BlobServiceClient
from docx import Document
from pptx import Presentation

from services.runtime_settings import blob_api_version, blob_connection_string, blob_container_name


def ingest_material_from_blob(
    course_id: str,
    material_id: str,
    blob_path: str,
    file_name: str,
    content_type: str,
) -> str:
    """Download the blob and extract plain text."""
    blob_bytes = _download_blob(blob_path)

    text = _extract_text(blob_bytes, file_name, content_type)

    normalized = "\n".join(line.strip() for line in text.splitlines() if line and line.strip())
    return normalized


def _download_blob(blob_path: str) -> bytes:
    connection_string = blob_connection_string()
    container_name = blob_container_name()

    client = BlobServiceClient.from_connection_string(
        connection_string,
        api_version=blob_api_version(),
    )
    blob_client = client.get_blob_client(container=container_name, blob=blob_path)
    downloader = blob_client.download_blob()
    return cast(bytes, downloader.readall())


def _extract_text(blob_bytes: bytes, file_name: str, content_type: str) -> str:
    lower_name = file_name.lower()
    lower_type = content_type.lower()

    if lower_name.endswith(".pdf") or lower_type == "application/pdf":
        try:
            with fitz.open(stream=blob_bytes, filetype="pdf") as doc:
                pages = [page.get_text("text") for page in doc]
            return "\n".join(page for page in pages if page and page.strip())
        except Exception:
            return ""

    if lower_name.endswith(".docx") or lower_type in {"application/vnd.openxmlformats-officedocument.wordprocessingml.document", "application/msword"}:
        try:
            document = Document(BytesIO(blob_bytes))
            paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text and paragraph.text.strip()]
            return "\n".join(paragraphs)
        except Exception:
            return ""

    if lower_name.endswith(".pptx") or lower_type == "application/vnd.openxmlformats-officedocument.presentationml.presentation":
        try:
            presentation = Presentation(BytesIO(blob_bytes))
            slides: list[str] = []
            for slide in presentation.slides:
                texts: list[str] = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        texts.append(shape.text)
                if texts:
                    slides.append("\n".join(texts))
            return "\n".join(slides)
        except Exception:
            return ""

    if lower_name.endswith(".md") or lower_name.endswith(".txt") or lower_type.startswith("text/"):
        for encoding in ("utf-8", "latin-1", "cp1250"):
            try:
                return blob_bytes.decode(encoding)
            except Exception:
                continue

    return ""

